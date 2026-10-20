import io
import time

import pytest
from openpyxl import Workbook
from sqlalchemy import select

from app.db.session import system_session
from app.models.beleg_zaehler import BelegZaehler
from app.models.kunde import Kunde
from tests.conftest import auth_headers, login

PW = "pw-123456"


async def _admin(client, make_mandant, make_user, name="Import"):
    mandant = await make_mandant(name)
    user = await make_user(mandant=mandant, role="mandant_admin", password=PW)
    token = await login(client, user.email, PW)
    return mandant, user, auth_headers(token)


def csv_bytes(text: str, kodierung="utf-8") -> bytes:
    return text.encode(kodierung)


async def analyse(client, kopf, entitaet, inhalt: bytes, name="daten.csv", blatt=None):
    daten = {"entitaet": entitaet}
    if blatt:
        daten["blatt"] = blatt
    r = await client.post("/api/import/analyse", headers=kopf, data=daten, files={"datei": (name, inhalt)})
    return r


def konfig(analyse_json, **optionen):
    mapping = {v["feld"]: v["spalte"] for v in analyse_json["vorschlag"]}
    return {"mapping": mapping, **optionen}


async def kompletter_lauf(client, kopf, entitaet, inhalt, name="daten.csv", **optionen):
    a = (await analyse(client, kopf, entitaet, inhalt, name)).json()
    k = konfig(a, **optionen)
    p = await client.post(f"/api/import/{a['import_lauf_id']}/probelauf", headers=kopf, json=k)
    return a, k, p


KUNDEN_CSV = (
    "Kd-Nr;Firma;Straße;PLZ;Ort;E-Mail;Telefon;USt-IdNr\n"
    "K-00007;Müller GmbH;Hauptstr. 1;01067;Dresden;info@mueller.de;0351 123;DE 123456789\n"
    ";Schmidt & Söhne;Weg 2;04109;Leipzig;;;\n"
    ";;Weg 3;04109;Leipzig;;;\n"
)


async def test_analyse_vorschlag_und_kodierungen(client, make_mandant, make_user):
    _, _, kopf = await _admin(client, make_mandant, make_user)
    for kodierung, name in (("utf-8", "UTF-8"), ("utf-8-sig", "UTF-8 (mit BOM)"), ("cp1252", "Windows-1252")):
        r = await analyse(client, kopf, "kunden", csv_bytes(KUNDEN_CSV, kodierung))
        assert r.status_code == 201, r.text
        j = r.json()
        assert j["kodierung"] == name
        assert j["trennzeichen"] == ";"
        assert j["zeilen_gesamt"] == 3
        vs = {v["feld"]: v for v in j["vorschlag"]}
        assert vs["kundennummer"]["spalte"] == 0 and vs["kundennummer"]["sicherheit"] == "sicher"
        assert vs["name"]["spalte"] == 1
        assert vs["plz"]["spalte"] == 3
        assert vs["ust_idnr"]["spalte"] == 7
        assert j["spalten"][1]["name"] == "Firma"
    # Umlaute bleiben korrekt
    assert "Müller GmbH" in j["spalten"][1]["beispiele"]


@pytest.mark.parametrize("trenner", [";", ",", "\t"])
async def test_trennzeichen_erkennung(client, make_mandant, make_user, trenner):
    _, _, kopf = await _admin(client, make_mandant, make_user)
    text = trenner.join(["Firma", "PLZ", "Ort"]) + "\n" + trenner.join(['"A, B"', "01067", "Dresden"]) + "\n"
    j = (await analyse(client, kopf, "kunden", csv_bytes(text))).json()
    assert j["trennzeichen"] == trenner
    assert [s["name"] for s in j["spalten"]] == ["Firma", "PLZ", "Ort"]


async def test_probelauf_und_import_kunden_konsistent(client, make_mandant, make_user):
    mandant, _, kopf = await _admin(client, make_mandant, make_user)
    a, k, p = await kompletter_lauf(client, kopf, "kunden", csv_bytes(KUNDEN_CSV))
    assert p.status_code == 200, p.text
    z = p.json()["zusammenfassung"]
    assert z["gesamt"] == 3 and z["fehler"] == 1 and z["anlegen"] == 2
    fehler = [x for x in p.json()["zeilen"] if x["status"] == "fehler"]
    assert fehler[0]["zeile"] == 4 and "Name" in fehler[0]["meldungen"][0]
    # Probelauf schreibt nichts
    async with system_session() as s:
        assert (await s.execute(select(Kunde).where(Kunde.mandant_id == mandant.id))).first() is None
    r = await client.post(f"/api/import/{a['import_lauf_id']}/ausfuehren", headers=kopf, json=k)
    assert r.status_code == 200, r.text
    assert r.json()["zusammenfassung"] == z
    assert r.json()["status"] == "abgeschlossen"
    async with system_session() as s:
        kunden = (await s.execute(select(Kunde).where(Kunde.mandant_id == mandant.id).order_by(Kunde.kundennummer))).scalars().all()
        nummern = [x.kundennummer for x in kunden]
        assert "K-00007" in nummern and len(kunden) == 2
        mueller = next(x for x in kunden if x.name == "Müller GmbH")
        assert mueller.ust_idnr == "DE123456789"
        assert mueller.adresse == {"strasse": "Hauptstr. 1", "plz": "01067", "ort": "Dresden"}
        assert mueller.ansprechpartner[0]["email"] == "info@mueller.de"
        # Nummernkreis nachgefuehrt: automatische Nummer kollidiert nicht, Zaehler > 7
        zaehler = await s.scalar(select(BelegZaehler.naechste_nummer).where(BelegZaehler.mandant_id == mandant.id, BelegZaehler.belegart == "kunde"))
        assert zaehler >= 9
    n = await client.post("/api/kunden", headers=kopf, json={"name": "Neu"})
    assert n.json()["kundennummer"] not in nummern
    # Fehlerbericht der Fehlerzeile
    fb = await client.get(f"/api/import/{a['import_lauf_id']}/fehlerbericht.csv", headers=kopf)
    assert fb.status_code == 200 and "Name/Firma fehlt" in fb.content.decode("utf-8-sig")
    # zweiter Lauf desselben Laufs geht nicht
    assert (await client.post(f"/api/import/{a['import_lauf_id']}/ausfuehren", headers=kopf, json=k)).status_code == 409


async def test_ausfuehren_braucht_probelauf_und_gleiche_konfig(client, make_mandant, make_user):
    _, _, kopf = await _admin(client, make_mandant, make_user)
    a = (await analyse(client, kopf, "kunden", csv_bytes(KUNDEN_CSV))).json()
    k = konfig(a)
    assert (await client.post(f"/api/import/{a['import_lauf_id']}/ausfuehren", headers=kopf, json=k)).status_code == 409
    await client.post(f"/api/import/{a['import_lauf_id']}/probelauf", headers=kopf, json=k)
    k2 = {**k, "modus": "alles_oder_nichts"}
    r = await client.post(f"/api/import/{a['import_lauf_id']}/ausfuehren", headers=kopf, json=k2)
    assert r.status_code == 409 and "Probelauf" in r.json()["detail"]


async def test_alles_oder_nichts(client, make_mandant, make_user, make_kunde):
    mandant, _, kopf = await _admin(client, make_mandant, make_user)
    a, k, _ = await kompletter_lauf(client, kopf, "kunden", csv_bytes(KUNDEN_CSV), modus="alles_oder_nichts")
    r = await client.post(f"/api/import/{a['import_lauf_id']}/ausfuehren", headers=kopf, json=k)
    assert r.json()["status"] == "abgebrochen"
    async with system_session() as s:
        assert (await s.execute(select(Kunde).where(Kunde.mandant_id == mandant.id))).first() is None


async def test_dubletten_aktionen(client, make_mandant, make_user, make_kunde):
    mandant, _, kopf = await _admin(client, make_mandant, make_user)
    await make_kunde(mandant=mandant, name="Müller GmbH", kundennummer="K-00001", adresse={"plz": "01067"})
    csvtext = "Kundennummer;Firma;PLZ;Ort\n;MUELLER GmbH;01067;Dresden\nK-00001;Anderer;99999;X\n;Neu AG;11111;Y\n;Neu AG;11111;Y\n"
    a, k, p = await kompletter_lauf(client, kopf, "kunden", csv_bytes(csvtext))
    zeilen = {x["zeile"]: x for x in p.json()["zeilen"]}
    assert zeilen[2]["status"] == "dublette" and zeilen[2]["dublette_art"] == "name_plz" and zeilen[2]["aktion"] == "ueberspringen"
    assert zeilen[3]["dublette_art"] == "kundennummer"
    assert zeilen[5]["dublette_von"] == "Zeile 4"
    # kundennummer: 'anlegen' ist nicht erlaubt
    k["dubletten_aktionen"] = {"kundennummer": "anlegen"}
    assert (await client.post(f"/api/import/{a['import_lauf_id']}/probelauf", headers=kopf, json=k)).status_code == 422
    # aktualisieren leerer Felder
    k["dubletten_aktionen"] = {"name_plz": "aktualisieren"}
    p = await client.post(f"/api/import/{a['import_lauf_id']}/probelauf", headers=kopf, json=k)
    assert {x["zeile"]: x for x in p.json()["zeilen"]}[2]["aktion"] == "aktualisieren"
    r = await client.post(f"/api/import/{a['import_lauf_id']}/ausfuehren", headers=kopf, json=k)
    assert r.json()["zusammenfassung"]["aktualisieren"] == 1
    alt = (await client.get("/api/kunden", headers=kopf, params={"q": "Müller"})).json()[0]
    assert alt["adresse"]["ort"] == "Dresden" and alt["name"] == "Müller GmbH"


async def test_validierung_leitweg_debitor_typ(client, make_mandant, make_user):
    _, _, kopf = await _admin(client, make_mandant, make_user)
    text = (
        "Firma;Kundentyp;Leitweg-ID;Debitorenkonto;USt-IdNr\n"
        "A;Behörde;04011000-1234512345-06;10001;XX\n"
        "B;gewerbe;123;;\n"
        "C;gewerbe;;12;\n"
        "D;Raumschiff;;10001;\n"
    )
    _, _, p = await kompletter_lauf(client, kopf, "kunden", csv_bytes(text))
    z = {x["zeile"]: x for x in p.json()["zeilen"]}
    assert z[2]["status"] == "warnung"  # Behoerde ist kein Synonym -> Typ unbekannt + USt-Format
    assert z[3]["status"] == "fehler" and "Leitweg" in z[3]["meldungen"][0]
    assert z[4]["status"] == "fehler" and "Debitorenkonto" in z[4]["meldungen"][0]
    assert z[5]["status"] == "fehler" and "Debitorenkonto" in z[5]["meldungen"][0]  # in der Datei doppelt


async def test_xlsx_datum_zahlen_fuehrende_nullen(client, make_mandant, make_user):
    mandant, _, kopf = await _admin(client, make_mandant, make_user)
    wb = Workbook()
    ws = wb.active
    ws.append(["Kundennummer", "Firma", "PLZ"])
    ws.append([10023, "Excel GmbH", 1067])
    ws["A2"].number_format = "000000"
    ws["C2"].number_format = "00000"
    puffer = io.BytesIO()
    wb.save(puffer)
    a, k, p = await kompletter_lauf(client, kopf, "kunden", puffer.getvalue(), name="k.xlsx")
    assert a["format"] == "xlsx" and a["blatt"]
    await client.post(f"/api/import/{a['import_lauf_id']}/ausfuehren", headers=kopf, json=k)
    async with system_session() as s:
        kunde = (await s.execute(select(Kunde).where(Kunde.mandant_id == mandant.id))).scalar_one()
        assert kunde.kundennummer == "010023" and kunde.adresse["plz"] == "01067"


async def test_xlsx_ablehnungen(client, make_mandant, make_user):
    _, _, kopf = await _admin(client, make_mandant, make_user)
    r = await analyse(client, kopf, "kunden", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"0" * 50, "alt.xls")
    assert r.status_code == 422 and "xls" in r.json()["detail"]
    r = await analyse(client, kopf, "kunden", b"", "leer.csv")
    assert r.status_code == 422
    r = await analyse(client, kopf, "kunden", b"a;b\n1;2\n", "x.exe")
    assert r.status_code == 422
    r = await analyse(client, kopf, "kunden", b"PK\x03\x04kaputt", "k.xlsx")
    assert r.status_code == 422


async def test_limits(client, make_mandant, make_user, monkeypatch):
    from app.core.config import get_settings

    _, _, kopf = await _admin(client, make_mandant, make_user)
    monkeypatch.setattr(get_settings(), "import_max_zeilen", 5)
    text = "Firma\n" + "\n".join(f"F{i}" for i in range(10))
    r = await analyse(client, kopf, "kunden", csv_bytes(text))
    assert r.status_code == 413
    monkeypatch.setattr(get_settings(), "import_max_bytes", 10)
    assert (await analyse(client, kopf, "kunden", csv_bytes("Firma\n" + "x" * 50))).status_code == 413


async def test_csv_injection_im_fehlerbericht(client, make_mandant, make_user):
    _, _, kopf = await _admin(client, make_mandant, make_user)
    text = "Firma;Ort;Notiz\n;=HYPERLINK(\"x\");+49 171\n"
    a, k, _ = await kompletter_lauf(client, kopf, "kunden", csv_bytes(text))
    fb = await client.get(f"/api/import/{a['import_lauf_id']}/fehlerbericht.csv", headers=kopf)
    inhalt = fb.content.decode("utf-8-sig")
    assert "'=HYPERLINK" in inhalt and "'+49 171" in inhalt


async def test_steuerzeichen_bereinigt(client, make_mandant, make_user):
    mandant, _, kopf = await _admin(client, make_mandant, make_user)
    a, k, _ = await kompletter_lauf(client, kopf, "kunden", csv_bytes("Firma;Ort\nA\x00B\x07C‮;Dre\x01sden\n"))
    await client.post(f"/api/import/{a['import_lauf_id']}/ausfuehren", headers=kopf, json=k)
    async with system_session() as s:
        kunde = (await s.execute(select(Kunde).where(Kunde.mandant_id == mandant.id))).scalar_one()
        assert kunde.name == "ABC" and kunde.adresse["ort"] == "Dresden"


async def test_rechte_und_mandantentrennung(client, make_mandant, make_user):
    m1, _, kopf1 = await _admin(client, make_mandant, make_user, "A")
    m2, _, kopf2 = await _admin(client, make_mandant, make_user, "B")
    a, k, _ = await kompletter_lauf(client, kopf1, "kunden", csv_bytes(KUNDEN_CSV))
    await client.post(f"/api/import/{a['import_lauf_id']}/ausfuehren", headers=kopf1, json=k)
    lid = a["import_lauf_id"]
    for methode, pfad in (("get", ""), ("get", "/fehlerbericht.csv"), ("post", "/rueckgaengig")):
        assert (await getattr(client, methode)(f"/api/import/{lid}{pfad}", headers=kopf2)).status_code == 404
    assert (await client.get("/api/import", headers=kopf2)).json() == []
    assert len((await client.get("/api/kunden", headers=kopf2)).json()) == 0
    # Techniker/Disponent haben keinen Zugriff
    tech = await make_user(mandant=m1, role="disponent", password=PW)
    kt = auth_headers(await login(client, tech.email, PW))
    assert (await client.get("/api/import", headers=kt)).status_code == 403
    assert (await analyse(client, kt, "kunden", csv_bytes(KUNDEN_CSV))).status_code == 403
    assert (await client.get("/api/import")).status_code == 401


async def test_modul_gesperrt(client, make_mandant, make_user):
    mandant, _, kopf = await _admin(client, make_mandant, make_user)
    async with system_session() as s:
        m = await s.get(type(mandant), mandant.id)
        m.deaktivierte_module = ["material"]
    assert (await analyse(client, kopf, "material", csv_bytes("Bezeichnung\nX\n"))).status_code == 403
    assert (await analyse(client, kopf, "kunden", csv_bytes("Firma\nX\n"))).status_code == 201


async def test_vorlagen(client, make_mandant, make_user):
    _, _, kopf = await _admin(client, make_mandant, make_user)
    for e in ("kunden", "anlagen", "material"):
        r = await client.get(f"/api/import/vorlagen/{e}.csv", headers=kopf)
        assert r.status_code == 200
        text = r.content.decode("utf-8-sig").splitlines()
        assert len(text) == 3
        # Vorlage ist importierbar und wird vollstaendig zugeordnet
        a = (await analyse(client, kopf, e, r.content)).json()
        assert a["zeilen_gesamt"] == 2
        assert sum(v["spalte"] is not None for v in a["vorschlag"]) == len(a["spalten"])
    assert (await client.get("/api/import/vorlagen/nix.csv", headers=kopf)).status_code == 404


async def test_performance_5000_kunden(client, make_mandant, make_user):
    mandant, _, kopf = await _admin(client, make_mandant, make_user)
    zeilen = ["Kundennummer;Firma;PLZ;Ort;E-Mail"] + [
        f"{'K-%05d' % i if i % 2 else ''};Firma {i};{10000 + i % 9000};Ort;f{i}@x.de" for i in range(1, 5001)
    ]
    a, k, p = await kompletter_lauf(client, kopf, "kunden", csv_bytes("\n".join(zeilen)))
    assert p.json()["zusammenfassung"]["fehler"] == 0
    t = time.monotonic()
    r = await client.post(f"/api/import/{a['import_lauf_id']}/ausfuehren", headers=kopf, json=k)
    dauer = time.monotonic() - t
    assert r.json()["zusammenfassung"]["anlegen"] == 5000
    assert dauer < 20, dauer
    async with system_session() as s:
        n = len((await s.execute(select(Kunde.id).where(Kunde.mandant_id == mandant.id))).all())
        assert n == 5000


async def test_rueckgaengig_kunden(client, make_mandant, make_user):
    mandant, _, kopf = await _admin(client, make_mandant, make_user)
    a, k, _ = await kompletter_lauf(client, kopf, "kunden", csv_bytes("Firma\nA\nB\nC\n"))
    await client.post(f"/api/import/{a['import_lauf_id']}/ausfuehren", headers=kopf, json=k)
    kunden = (await client.get("/api/kunden", headers=kopf)).json()
    assert len(kunden) == 3
    # B wird geaendert, C bekommt einen Auftrag
    b = next(x for x in kunden if x["name"] == "B")
    await client.patch(f"/api/kunden/{b['id']}", headers=kopf, json={"notiz": "x"})
    c = next(x for x in kunden if x["name"] == "C")
    v = await client.post("/api/vorgaenge", headers=kopf, json={"kunde_id": c["id"], "titel": "T", "abrechnungsart": "aufwand", "leistungstyp": "stoerung"})
    assert v.status_code == 201, v.text
    r = await client.post(f"/api/import/{a['import_lauf_id']}/rueckgaengig", headers=kopf)
    assert r.status_code == 200, r.text
    ergebnis = r.json()["ergebnis"]["rueckgaengig"]
    assert ergebnis["zurueckgesetzt"] == {"kunde": 1} and ergebnis["uebersprungen_gesamt"] == 2
    assert [x["name"] for x in (await client.get("/api/kunden", headers=kopf)).json()] == ["B", "C"]
    assert r.json()["status"] == "rueckgaengig"
    assert (await client.post(f"/api/import/{a['import_lauf_id']}/rueckgaengig", headers=kopf)).status_code == 409


async def test_laufliste_und_verwerfen(client, make_mandant, make_user):
    _, _, kopf = await _admin(client, make_mandant, make_user)
    a = (await analyse(client, kopf, "kunden", csv_bytes("Firma\nA\n"))).json()
    liste = (await client.get("/api/import", headers=kopf)).json()
    assert liste[0]["id"] == a["import_lauf_id"] and liste[0]["quelle_verfuegbar"] and len(liste[0]["sha256"]) == 64
    assert (await client.delete(f"/api/import/{a['import_lauf_id']}", headers=kopf)).status_code == 204
    g = (await client.get(f"/api/import/{a['import_lauf_id']}", headers=kopf)).json()
    assert g["status"] == "abgelaufen" and not g["quelle_verfuegbar"]
    r = await client.post(f"/api/import/{a['import_lauf_id']}/probelauf", headers=kopf, json={"mapping": {"name": 0}})
    assert r.status_code == 409


async def test_aufraeumen(client, make_mandant, make_user):
    from datetime import UTC, datetime, timedelta

    from app.services.datenimport.lauf import raeume_auf

    _, _, kopf = await _admin(client, make_mandant, make_user)
    a = (await analyse(client, kopf, "kunden", csv_bytes("Firma\nA\n"))).json()
    assert (await raeume_auf())["quellen"] == 0
    erg = await raeume_auf(datetime.now(UTC) + timedelta(hours=25))
    assert erg["quellen"] == 1
    assert (await client.get(f"/api/import/{a['import_lauf_id']}", headers=kopf)).json()["status"] == "abgelaufen"


async def test_migration_up_down_up():
    import os

    from alembic import command
    from alembic.config import Config

    from tests.conftest import _BACKEND_DIR

    cfg = Config(os.path.join(_BACKEND_DIR, "alembic.ini"))
    cfg.set_main_option("script_location", os.path.join(_BACKEND_DIR, "alembic"))
    from sqlalchemy import create_engine, inspect

    from app.core.config import get_settings

    engine = create_engine(get_settings().database_url_sync)
    try:
        command.downgrade(cfg, "0111")
        assert "import_laeufe" not in inspect(engine).get_table_names()
        command.upgrade(cfg, "head")
        assert {"import_laeufe", "import_lauf_eintraege"} <= set(inspect(engine).get_table_names())
    finally:
        engine.dispose()
