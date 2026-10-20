"""Liest CSV- und XLSX-Dateien des Datenimports in eine einheitliche Tabelle.

Die Datei ist nicht vertrauenswuerdig: Groesse, Zeilen-, Spalten- und Zellenlaenge sind
begrenzt, XLSX wird nur gelesen (read_only, data_only -- Formeln liefern ihren
gespeicherten Wert), Makros, externe Verknuepfungen und XML-Entities werden abgelehnt.
Der Dateiinhalt landet in keiner Fehlermeldung."""
import csv
import io
import re
import time
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime, time as tageszeit
from decimal import Decimal

from app.services.datenimport.werte import ExcelZahl, bereinige_text, datum_als_text

_CSV_ENDUNGEN = {"csv", "txt", "tsv"}
_ZIP_MAGIC = b"PK\x03\x04"
_OLE_MAGIC = bytes.fromhex("D0CF11E0A1B11AE1")
_KANDIDATEN = (";", "\t", ",", "|")
_ERKENNUNG_ZEILEN = 50


class ImportDateiFehler(ValueError):
    """Datei nicht lesbar oder ueber den Limits -- die Meldung ist fuer den Nutzer gedacht."""

    def __init__(self, meldung: str, *, status_code: int = 422):
        super().__init__(meldung)
        self.status_code = status_code


@dataclass(frozen=True)
class Limits:
    max_bytes: int
    max_zeilen: int
    max_spalten: int
    max_zellenlaenge: int
    max_unkomprimiert: int
    zeitlimit_sekunden: float


@dataclass
class Tabelle:
    format: str
    kopf: list[str]
    zeilen: list[list[str]]
    # Zeilennummer in der Datei (1 = Kopfzeile), damit Meldungen zur Tabellenkalkulation passen.
    zeilennummern: list[int]
    trennzeichen: str | None = None
    kodierung: str | None = None
    blatt: str | None = None
    blaetter: list[str] = field(default_factory=list)


def lese_tabelle(daten: bytes, dateiname: str, limits: Limits, *, blatt: str | None = None) -> Tabelle:
    if not daten:
        raise ImportDateiFehler("Die Datei ist leer.")
    if len(daten) > limits.max_bytes:
        raise ImportDateiFehler(
            f"Die Datei ist größer als {limits.max_bytes // (1024 * 1024)} MB.", status_code=413
        )
    endung = dateiname.rsplit(".", 1)[-1].lower() if "." in dateiname else ""
    if endung in {"xls", "xlsm", "xlsb", "xltm"} or daten.startswith(_OLE_MAGIC):
        raise ImportDateiFehler(
            "Dieses Excel-Format wird nicht unterstützt (alte .xls-Dateien und Dateien mit Makros). "
            "Bitte als .xlsx oder CSV speichern."
        )
    frist = time.monotonic() + limits.zeitlimit_sekunden
    if daten.startswith(_ZIP_MAGIC):
        if endung not in {"xlsx", ""}:
            raise ImportDateiFehler("Dateiendung und Inhalt passen nicht zusammen (erwartet .xlsx oder .csv).")
        return _lese_xlsx(daten, limits, blatt, frist)
    if endung not in _CSV_ENDUNGEN:
        raise ImportDateiFehler("Es werden nur CSV- und XLSX-Dateien unterstützt.")
    return _lese_csv(daten, limits, frist)


# --- CSV -------------------------------------------------------------------


def _dekodieren(daten: bytes) -> tuple[str, str]:
    if daten.startswith((b"\xff\xfe", b"\xfe\xff")):
        try:
            return daten.decode("utf-16"), "UTF-16"
        except UnicodeDecodeError:
            raise ImportDateiFehler("Die Datei konnte nicht gelesen werden (Kodierung).") from None
    if b"\x00" in daten[:4096]:
        raise ImportDateiFehler("Die Datei ist keine Textdatei (CSV).")
    if daten.startswith(b"\xef\xbb\xbf"):
        try:
            return daten[3:].decode("utf-8"), "UTF-8 (mit BOM)"
        except UnicodeDecodeError:
            pass
    else:
        try:
            return daten.decode("utf-8"), "UTF-8"
        except UnicodeDecodeError:
            pass
    try:
        return daten.decode("cp1252"), "Windows-1252"
    except UnicodeDecodeError:
        return daten.decode("latin-1"), "ISO-8859-1"


def _trennzeichen_erkennen(text: str) -> str:
    erste = text.split("\n", 1)[0].strip()
    if re.fullmatch(r"sep=.", erste, flags=re.IGNORECASE):
        return erste[-1]
    probe = io.StringIO(text[:200_000])
    beste, beste_wertung = ";", (0, 0.0)
    for kandidat in _KANDIDATEN:
        probe.seek(0)
        try:
            zeilen = []
            for zeile in csv.reader(probe, delimiter=kandidat):
                if zeile:
                    zeilen.append(len(zeile))
                if len(zeilen) >= _ERKENNUNG_ZEILEN:
                    break
        except csv.Error:
            continue
        if not zeilen:
            continue
        haeufigste = max(set(zeilen), key=zeilen.count)
        if haeufigste < 2:
            continue
        # Wertung: Spaltenzahl der haeufigsten Zeilenlaenge, gewichtet mit ihrem Anteil;
        # bei Gleichstand gewinnt die frueher stehende Variante (";" vor Tab vor ",").
        anteil = zeilen.count(haeufigste) / len(zeilen)
        wertung = (1 if anteil >= 0.6 else 0, anteil * haeufigste)
        if wertung > beste_wertung:
            beste, beste_wertung = kandidat, wertung
    return beste


def _lese_csv(daten: bytes, limits: Limits, frist: float) -> Tabelle:
    text, kodierung = _dekodieren(daten)
    trennzeichen = _trennzeichen_erkennen(text)
    if re.match(r"(?i)sep=.\r?\n", text):
        text = text.split("\n", 1)[1]
    leser = csv.reader(io.StringIO(text, newline=""), delimiter=trennzeichen)
    zeilen_roh: list[tuple[int, list[str]]] = []
    try:
        for nr, zeile in enumerate(leser, start=1):
            if nr % 1000 == 0 and time.monotonic() > frist:
                raise ImportDateiFehler("Die Datei konnte nicht rechtzeitig gelesen werden (Zeitlimit).", status_code=408)
            if len(zeile) > limits.max_spalten:
                raise ImportDateiFehler(f"Die Datei hat mehr als {limits.max_spalten} Spalten.")
            zellen = [_zelle_text(z, limits, nr) for z in zeile]
            if any(zellen):
                zeilen_roh.append((nr, zellen))
            if len(zeilen_roh) > limits.max_zeilen + 1:
                raise _zu_viele_zeilen(limits)
    except csv.Error:
        raise ImportDateiFehler("Die CSV-Datei ist fehlerhaft aufgebaut (z. B. nicht geschlossene Anführungszeichen).") from None
    return _zusammenbauen("csv", zeilen_roh, trennzeichen=trennzeichen, kodierung=kodierung)


def _zelle_text(wert: str, limits: Limits, zeile: int) -> str:
    if len(wert) > limits.max_zellenlaenge:
        raise ImportDateiFehler(
            f"Zeile {zeile}: Eine Zelle ist länger als {limits.max_zellenlaenge} Zeichen – "
            "die Datei ist vermutlich keine Tabelle im erwarteten Format."
        )
    return bereinige_text(wert, einzeilig=False)


def _zu_viele_zeilen(limits: Limits) -> ImportDateiFehler:
    return ImportDateiFehler(
        f"Die Datei hat mehr als {limits.max_zeilen:,} Datenzeilen. Bitte in mehrere Dateien aufteilen.".replace(",", "."),
        status_code=413,
    )


# --- XLSX ------------------------------------------------------------------


def _xlsx_pruefen(daten: bytes, limits: Limits) -> zipfile.ZipFile:
    try:
        archiv = zipfile.ZipFile(io.BytesIO(daten))
        eintraege = archiv.infolist()
    except (zipfile.BadZipFile, OSError):
        raise ImportDateiFehler("Die XLSX-Datei ist beschädigt oder keine Excel-Datei.") from None
    if len(eintraege) > 2000:
        raise ImportDateiFehler("Die XLSX-Datei enthält unplausibel viele Bestandteile.")
    gesamt = sum(e.file_size for e in eintraege)
    if gesamt > limits.max_unkomprimiert:
        raise ImportDateiFehler(
            f"Die XLSX-Datei ist entpackt größer als {limits.max_unkomprimiert // (1024 * 1024)} MB.", status_code=413
        )
    namen = {e.filename for e in eintraege}
    if "xl/workbook.xml" not in namen:
        raise ImportDateiFehler("Die Datei ist keine gültige Excel-Arbeitsmappe (.xlsx).")
    if any(n.lower().endswith("vbaproject.bin") for n in namen):
        raise ImportDateiFehler("Excel-Dateien mit Makros werden nicht importiert. Bitte als .xlsx ohne Makros speichern.")
    if any(n.startswith("xl/externalLinks/") for n in namen):
        raise ImportDateiFehler(
            "Die Arbeitsmappe enthält Verknüpfungen zu anderen Dateien. Bitte die Werte kopieren und "
            "in eine neue Arbeitsmappe einfügen."
        )
    for eintrag in eintraege:
        if eintrag.filename.lower().endswith((".xml", ".rels")) and eintrag.file_size:
            with archiv.open(eintrag) as teil:
                anfang = teil.read(4096)
            if b"<!DOCTYPE" in anfang or b"<!ENTITY" in anfang:
                raise ImportDateiFehler("Die Excel-Datei enthält unzulässige XML-Definitionen.")
    return archiv


def _lese_xlsx(daten: bytes, limits: Limits, blatt: str | None, frist: float) -> Tabelle:
    from openpyxl import load_workbook

    _xlsx_pruefen(daten, limits).close()
    try:
        mappe = load_workbook(io.BytesIO(daten), read_only=True, data_only=True, keep_links=False)
    except Exception:
        raise ImportDateiFehler("Die Excel-Datei konnte nicht gelesen werden.") from None
    try:
        blaetter = [name for name in mappe.sheetnames if mappe[name].sheet_state == "visible"] or list(mappe.sheetnames)
        if not blaetter:
            raise ImportDateiFehler("Die Arbeitsmappe enthält kein Tabellenblatt.")
        if blatt is not None and blatt not in blaetter:
            raise ImportDateiFehler("Das gewählte Tabellenblatt existiert nicht in der Datei.")
        gewaehlt = blatt or blaetter[0]
        blatt_objekt = mappe[gewaehlt]
        zeilen_roh: list[tuple[int, list[str]]] = []
        try:
            # Manche Programme schreiben eine falsche Bereichsangabe (A1:A1): ohne Zuruecksetzen
            # wuerden Zeilen abgeschnitten.
            blatt_objekt.reset_dimensions()
            for nr, zeile in enumerate(blatt_objekt.iter_rows(), start=1):
                if nr % 500 == 0 and time.monotonic() > frist:
                    raise ImportDateiFehler(
                        "Die Datei konnte nicht rechtzeitig gelesen werden (Zeitlimit).", status_code=408
                    )
                if len(zeile) > limits.max_spalten:
                    # Formatierte Leerspalten bis XFD sind ueblich: erst kuerzen, dann pruefen.
                    zeile = _ohne_leere_enden(zeile)
                    if len(zeile) > limits.max_spalten:
                        raise ImportDateiFehler(f"Das Tabellenblatt hat mehr als {limits.max_spalten} Spalten.")
                zellen = [_xlsx_zelle(z, limits, nr) for z in zeile]
                if any(zellen):
                    zeilen_roh.append((nr, zellen))
                if len(zeilen_roh) > limits.max_zeilen + 1:
                    raise _zu_viele_zeilen(limits)
        except ImportDateiFehler:
            raise
        except Exception:
            raise ImportDateiFehler("Die Excel-Datei konnte nicht gelesen werden.") from None
    finally:
        mappe.close()
    return _zusammenbauen("xlsx", zeilen_roh, blatt=gewaehlt, blaetter=blaetter)


def _ohne_leere_enden(zeile) -> list:
    ende = len(zeile)
    while ende and getattr(zeile[ende - 1], "value", None) in (None, ""):
        ende -= 1
    return list(zeile[:ende])


def _xlsx_zelle(zelle, limits: Limits, zeile: int) -> str:
    wert = getattr(zelle, "value", None)
    if wert is None:
        return ""
    if isinstance(wert, bool):
        return "Ja" if wert else "Nein"
    if isinstance(wert, (datetime, date)):
        return datum_als_text(wert)
    if isinstance(wert, tageszeit):
        return wert.replace(microsecond=0).isoformat()
    if isinstance(wert, (int, float, Decimal)):
        return _zahl_zelle(wert, getattr(zelle, "number_format", "General") or "General")
    return _zelle_text(str(wert), limits, zeile)


_NULLFORMAT = re.compile(r"^0{2,}$")


def _zahl_zelle(wert: int | float | Decimal, zahlenformat: str) -> ExcelZahl:
    """Excel speichert PLZ/Kundennummern oft als Zahl mit Format "00000": die fuehrenden
    Nullen stehen nur im Format, nicht im Wert."""
    zahl = Decimal(f"{wert:.15g}") if isinstance(wert, float) else Decimal(wert)
    if not zahl.is_finite():
        return ExcelZahl("", Decimal(0))
    ganzzahlig = zahl == zahl.to_integral_value()
    if ganzzahlig:
        text = str(int(zahl))
        if _NULLFORMAT.match(zahlenformat) and int(zahl) >= 0:
            text = text.zfill(len(zahlenformat))
    else:
        text = format(zahl.normalize(), "f")
    return ExcelZahl(text, zahl)


# --- gemeinsam ---------------------------------------------------------------


def _zusammenbauen(
    format_: str,
    zeilen_roh: list[tuple[int, list[str]]],
    *,
    trennzeichen: str | None = None,
    kodierung: str | None = None,
    blatt: str | None = None,
    blaetter: list[str] | None = None,
) -> Tabelle:
    if not zeilen_roh:
        raise ImportDateiFehler("Die Datei enthält keine Daten.")
    (_, kopf_roh), daten_zeilen = zeilen_roh[0], zeilen_roh[1:]
    if not daten_zeilen:
        raise ImportDateiFehler("Die Datei enthält nur eine Kopfzeile, aber keine Datenzeilen.")
    breite = max(len(kopf_roh), *(len(z) for _, z in daten_zeilen))
    # Spalten ohne Kopf und ohne jeden Wert (Formatierungsreste) fallen weg.
    benutzt = [
        i for i in range(breite)
        if (i < len(kopf_roh) and kopf_roh[i]) or any(i < len(z) and z[i] for _, z in daten_zeilen)
    ]
    kopf: list[str] = []
    gesehen: dict[str, int] = {}
    for i in benutzt:
        name = bereinige_text(kopf_roh[i]) if i < len(kopf_roh) else ""
        name = name[:120] or f"Spalte {i + 1}"
        gesehen[name] = gesehen.get(name, 0) + 1
        kopf.append(name if gesehen[name] == 1 else f"{name} ({gesehen[name]})")
    zeilen = [[z[i] if i < len(z) else "" for i in benutzt] for _, z in daten_zeilen]
    return Tabelle(
        format=format_,
        kopf=kopf,
        zeilen=zeilen,
        zeilennummern=[nr for nr, _ in daten_zeilen],
        trennzeichen=trennzeichen,
        kodierung=kodierung,
        blatt=blatt,
        blaetter=blaetter or [],
    )
